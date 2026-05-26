from PyQt6.QtWidgets import (QWidget, QVBoxLayout, QPushButton,
                             QTableWidget, QTableWidgetItem, QLabel, QMessageBox,
                             QHBoxLayout, QDateEdit, QComboBox, QHeaderView, QScrollArea)
from PyQt6.QtCore import QDate, Qt
from api_client import APIClient
from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import os
from datetime import datetime

DESKTOP = os.path.join(os.path.expanduser("~"), "Desktop")
REPORTS_DIR = os.path.join(DESKTOP, "Отчеты")

STATUS_RUSSIAN = {
    "accepted": "Принят",
    "classified": "Классифицирован",
    "processing": "В обработке",
    "completed": "Завершен",
    "registered": "Зарегистрирован",
    "verified": "Проверен",
    "rejected": "Отклонен"
}


class ReportTab(QWidget):
    def __init__(self, user):
        super().__init__()
        self.user = user
        self.client = APIClient()
        self.init_ui()
        self.load_data()

    def init_ui(self):
        layout = QVBoxLayout()
        layout.setContentsMargins(20, 20, 20, 20)
        layout.setSpacing(16)

        title = QLabel("Отчеты и аналитика")
        title.setObjectName("title")
        layout.addWidget(title)

        # Фильтры
        filter_layout = QHBoxLayout()
        filter_layout.addWidget(QLabel("Дата от:"))
        self.date_from = QDateEdit()
        self.date_from.setDate(QDate(2026, 1, 1))
        self.date_from.setCalendarPopup(True)
        filter_layout.addWidget(self.date_from)

        filter_layout.addWidget(QLabel("Дата до:"))
        self.date_to = QDateEdit()
        self.date_to.setDate(QDate.currentDate())
        self.date_to.setCalendarPopup(True)
        filter_layout.addWidget(self.date_to)

        filter_layout.addWidget(QLabel("Статус:"))
        self.status_combo = QComboBox()
        self.status_combo.addItems(["Все", "Принят", "Классифицирован", "Завершен"])
        self.status_combo.currentTextChanged.connect(self.load_data)
        filter_layout.addWidget(self.status_combo)

        filter_layout.addStretch()
        layout.addLayout(filter_layout)

        # Скролл для таблицы
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setStyleSheet("QScrollArea { border: none; background-color: transparent; }")

        self.table = QTableWidget()
        self.table.setColumnCount(8)
        self.table.setHorizontalHeaderLabels(
            ["Код", "Наименование", "Класс", "Объем (т)", "ФККО", "Дата", "Статус", "Цех"])
        self.table.setAlternatingRowColors(True)
        self.table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self.table.setWordWrap(True)
        self.table.horizontalHeader().setStretchLastSection(True)

        # Стили для таблицы
        self.table.setStyleSheet("""
            QTableWidget {
                font-size: 12pt;
                background-color: #ffffff;
                alternate-background-color: #f8f9fa;
                gridline-color: #e0e4e8;
            }
            QTableWidget::item {
                padding: 10px;
            }
            QHeaderView::section {
                background-color: #f8f9fa;
                color: #495057;
                font-size: 13pt;
                font-weight: bold;
                padding: 12px;
            }
        """)

        scroll.setWidget(self.table)
        layout.addWidget(scroll)

        # Кнопки
        btn_layout = QHBoxLayout()
        self.refresh_btn = QPushButton("🔄 Обновить")
        self.refresh_btn.setMinimumHeight(40)
        self.refresh_btn.clicked.connect(self.load_data)

        self.excel_btn = QPushButton("📊 Экспорт в Excel")
        self.excel_btn.setMinimumHeight(40)
        self.excel_btn.clicked.connect(self.export_excel)

        self.word_btn = QPushButton("📄 Экспорт выбранной партии в Word")
        self.word_btn.setMinimumHeight(40)
        self.word_btn.clicked.connect(self.export_word)

        btn_layout.addWidget(self.refresh_btn)
        btn_layout.addWidget(self.excel_btn)
        btn_layout.addWidget(self.word_btn)
        layout.addLayout(btn_layout)

        # Информационная метка
        self.info_label = QLabel("")
        self.info_label.setStyleSheet("color: #6c757d; font-size: 11pt;")
        layout.addWidget(self.info_label)

        self.setLayout(layout)

    def ensure_reports_dir(self):
        if not os.path.exists(REPORTS_DIR):
            os.makedirs(REPORTS_DIR)
        return REPORTS_DIR

    def load_data(self):
        try:
            batches = self.client.get_batches()
            print(f"Загружено партий: {len(batches)}")

            date_from = self.date_from.date().toString("yyyy-MM-dd")
            date_to = self.date_to.date().toString("yyyy-MM-dd")
            status_ru_filter = self.status_combo.currentText()

            # Преобразуем русский статус в английский для фильтрации
            status_en_filter = None
            for en, ru in STATUS_RUSSIAN.items():
                if ru == status_ru_filter:
                    status_en_filter = en
                    break

            filtered = []
            for batch in batches:
                received = batch.get('received_at', '')[:10] if batch.get('received_at') else ''
                if received < date_from or received > date_to:
                    continue
                if status_en_filter and batch.get('status') != status_en_filter:
                    continue
                filtered.append(batch)

            self.table.setRowCount(len(filtered))
            self.batches_data = filtered

            for i, batch in enumerate(filtered):
                self.table.setItem(i, 0, QTableWidgetItem(batch.get('code', '')))
                self.table.setItem(i, 1, QTableWidgetItem(batch.get('name', '')[:50] if batch.get('name') else ''))

                hazard = batch.get('hazard_class', '')
                hazard_text = {1: "I", 2: "II", 3: "III", 4: "IV", 5: "V"}.get(hazard, str(hazard))
                self.table.setItem(i, 2, QTableWidgetItem(hazard_text))

                vol = batch.get('volume_tons', 0)
                self.table.setItem(i, 3, QTableWidgetItem(f"{vol:.2f}" if vol else "0.00"))
                self.table.setItem(i, 4, QTableWidgetItem(batch.get('fkko_code', '—')[:25]))
                self.table.setItem(i, 5, QTableWidgetItem(
                    batch.get('received_at', '')[:10] if batch.get('received_at') else '—'))

                status_ru = STATUS_RUSSIAN.get(batch.get('status', '—'), batch.get('status', '—'))
                self.table.setItem(i, 6, QTableWidgetItem(status_ru))
                self.table.setItem(i, 7, QTableWidgetItem(batch.get('source_department', '—') or '—'))

            # Устанавливаем ширину колонок
            self.table.setColumnWidth(0, 100)
            self.table.setColumnWidth(1, 300)
            self.table.setColumnWidth(2, 70)
            self.table.setColumnWidth(3, 90)
            self.table.setColumnWidth(4, 160)
            self.table.setColumnWidth(5, 110)
            self.table.setColumnWidth(6, 120)

            self.table.resizeRowsToContents()
            self.info_label.setText(f"📊 Найдено записей: {len(filtered)}")

        except Exception as e:
            print(f"Ошибка загрузки: {e}")
            self.info_label.setText(f"❌ Ошибка загрузки: {e}")

    def export_excel(self):
        if not hasattr(self, 'batches_data') or not self.batches_data:
            QMessageBox.warning(self, "Ошибка", "Нет данных для экспорта")
            return

        try:
            reports_dir = self.ensure_reports_dir()
            wb = Workbook()
            ws = wb.active
            ws.title = "Отчет по партиям"

            headers = ["Код партии", "Наименование отхода", "Класс опасности",
                       "Объем (тонн)", "Код ФККО", "Дата поступления", "Статус", "Цех-источник"]

            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="0066cc", end_color="0066cc", fill_type="solid")

            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center")

            for row, batch in enumerate(self.batches_data, 2):
                ws.cell(row=row, column=1, value=batch.get('code', ''))
                ws.cell(row=row, column=2, value=batch.get('name', ''))

                hazard = batch.get('hazard_class', '')
                hazard_text = {1: "I", 2: "II", 3: "III", 4: "IV", 5: "V"}.get(hazard, str(hazard))
                ws.cell(row=row, column=3, value=hazard_text)

                ws.cell(row=row, column=4, value=batch.get('volume_tons', 0))
                ws.cell(row=row, column=5, value=batch.get('fkko_code', ''))
                ws.cell(row=row, column=6, value=batch.get('received_at', '')[:10] if batch.get('received_at') else '')

                status_ru = STATUS_RUSSIAN.get(batch.get('status', ''), batch.get('status', ''))
                ws.cell(row=row, column=7, value=status_ru)
                ws.cell(row=row, column=8, value=batch.get('source_department', ''))

            for col in range(1, 9):
                ws.column_dimensions[chr(64 + col)].width = 20

            filename = os.path.join(reports_dir, f"report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
            wb.save(filename)
            QMessageBox.information(self, "Успех", f"Отчет сохранен:\n{filename}")

        except Exception as e:
            QMessageBox.critical(self, "Ошибка", f"Не удалось создать Excel: {e}")

    def export_word(self):
        selected = self.table.currentRow()
        if selected < 0 or not hasattr(self, 'batches_data') or selected >= len(self.batches_data):
            QMessageBox.warning(self, "Ошибка", "Выберите партию из таблицы (кликните на строку)")
            return

        batch = self.batches_data[selected]

        try:
            reports_dir = self.ensure_reports_dir()
            doc = Document()

            title = doc.add_heading(f"АКТ ПРИЕМА-ПЕРЕДАЧИ ОТХОДОВ №{batch.get('code', '')}", 0)
            title.alignment = 1

            doc.add_paragraph(f"Дата составления: {datetime.now().strftime('%d.%m.%Y')}")
            doc.add_paragraph()

            doc.add_heading("1. Сведения о партии отходов", level=1)
            doc.add_paragraph(f"Код партии: {batch.get('code', '')}")
            doc.add_paragraph(f"Наименование отхода: {batch.get('name', '')}")
            doc.add_paragraph(f"Код ФККО: {batch.get('fkko_code', '—')}")

            hazard = batch.get('hazard_class', '')
            hazard_text = {1: "I (чрезвычайно опасный)", 2: "II (высокоопасный)",
                           3: "III (умеренно опасный)", 4: "IV (малоопасный)",
                           5: "V (практически неопасный)"}.get(hazard, str(hazard))
            doc.add_paragraph(f"Класс опасности: {hazard_text}")

            doc.add_paragraph(f"Объем: {batch.get('volume_tons', 0)} тонн")
            doc.add_paragraph(
                f"Дата поступления: {batch.get('received_at', '')[:10] if batch.get('received_at') else '—'}")
            doc.add_paragraph(f"Цех-источник: {batch.get('source_department', '—')}")

            status_ru = STATUS_RUSSIAN.get(batch.get('status', ''), batch.get('status', ''))
            doc.add_paragraph(f"Статус: {status_ru}")

            doc.add_paragraph()
            doc.add_paragraph("От оператора склада: ___________________")
            doc.add_paragraph("От цеха-источника: ___________________")

            filename = os.path.join(reports_dir,
                                    f"act_{batch.get('code', '')}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx")
            doc.save(filename)
            QMessageBox.information(self, "Успех", f"Акт сохранен:\n{filename}")

        except Exception as e:
            QMessageBox.critical(self, "Ошибка", f"Не удалось создать Word: {e}")